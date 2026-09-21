"""Multi-format document ingestion and verification module.

Supports parsing plain text (.txt), PDF (.pdf), Word (.docx), and Excel (.xlsx)
attachments, with defensive screening for corrupt/unreadable files and non-BL
document types (Commercial Invoices, Packing Lists, Certificates of Origin).
"""

from pathlib import Path
from typing import Tuple, Optional
import pypdf
import docx
import openpyxl


_PARSE_CACHE: Dict[Tuple[str, float, int], Tuple[Optional[str], Optional[str]]] = {}

def parse_document(file_path: Path) -> Tuple[Optional[str], Optional[str]]:
    """Safely extract readable text from a multi-format attachment with in-memory caching.

    Args:
        file_path: Path to the attachment file.

    Returns:
        Tuple of (extracted_text, error_reason).
        error_reason is None on success, or 'missing' / 'unreadable' on failure.
    """
    if not file_path.exists():
        return None, "missing"
    
    st_stat = file_path.stat()
    if st_stat.st_size == 0:
        return None, "unreadable"

    cache_key = (str(file_path.resolve()), st_stat.st_mtime, st_stat.st_size)
    if cache_key in _PARSE_CACHE:
        return _PARSE_CACHE[cache_key]
    
    ext = file_path.suffix.lower()
    try:
        if ext == ".txt":
            res = (file_path.read_text(encoding="utf-8", errors="replace"), None)
            _PARSE_CACHE[cache_key] = res
            return res
        elif ext == ".pdf":
            reader = pypdf.PdfReader(file_path)
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
            if len(text.strip()) < 20 or "NO OCR TEXT LAYER" in text:
                res = (None, "unreadable")
            else:
                res = (text, None)
            _PARSE_CACHE[cache_key] = res
            return res
        elif ext == ".docx":
            doc = docx.Document(file_path)
            lines = [p.text for p in doc.paragraphs if p.text]
            for t in doc.tables:
                for row in t.rows:
                    cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                    lines.append(" : ".join(cells))
            res = ("\n".join(lines), None)
            _PARSE_CACHE[cache_key] = res
            return res
        elif ext == ".xlsx":
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            lines = []
            for name in wb.sheetnames:
                ws = wb[name]
                for row in ws.iter_rows(values_only=True):
                    row_vals = [str(c) for c in row if c is not None]
                    if row_vals:
                        lines.append(" : ".join(row_vals))
            wb.close()
            res = ("\n".join(lines), None)
            _PARSE_CACHE[cache_key] = res
            return res
    except Exception:
        res = (None, "unreadable")
        _PARSE_CACHE[cache_key] = res
        return res
    res = (None, "unreadable")
    _PARSE_CACHE[cache_key] = res
    return res

def check_wrong_doc_type(text: str) -> bool:
    """Detect if an attachment is an invalid document type (invoice, packing list, COO).

    Args:
        text: Raw extracted document text.

    Returns:
        True if the document is an invalid type that requires human review.
    """
    t_upper = text.upper()
    if "COMMERCIAL INVOICE" in t_upper and "BILL OF LADING" not in t_upper:
        return True
    if "PACKING LIST" in t_upper and "BILL OF LADING" not in t_upper:
        return True
    if "CERTIFICATE OF ORIGIN" in t_upper and "BILL OF LADING" not in t_upper:
        return True
    return False


